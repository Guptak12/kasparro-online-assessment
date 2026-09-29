from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .config import Settings
from .models import ParsedResume, ResumeFile


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ResumeParseError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True, slots=True)
class DiscoveryResult:
    files: list[Path]
    unsupported: list[Path]

    @property
    def total_discovered(self) -> int:
        return len(self.files) + len(self.unsupported)


def discover_resumes(directory: Path, recursive: bool = False) -> DiscoveryResult:
    root = directory.expanduser().resolve()
    if not root.is_dir():
        raise ValueError(f"Input directory does not exist or is not a directory: {root}")
    iterator = root.rglob("*") if recursive else root.glob("*")
    files: list[Path] = []
    unsupported: list[Path] = []
    for path in sorted(iterator, key=lambda item: str(item).lower()):
        if not path.is_file() or path.is_symlink() or path.name.startswith("."):
            continue
        if path.suffix.lower() in SUPPORTED_EXTENSIONS:
            files.append(path)
        else:
            unsupported.append(path)
    return DiscoveryResult(files=files, unsupported=unsupported)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalise_text(text: str, limit: int) -> tuple[str, bool]:
    text = unicodedata.normalize("NFKC", text)
    text = "".join(ch for ch in text if ch in "\n\t" or not unicodedata.category(ch).startswith("C"))
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = text.strip()
    truncated = len(text) > limit
    return text[:limit], truncated


def _extract_pdf(path: Path) -> tuple[str, list[str], str]:
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ResumeParseError("encrypted_pdf", "Encrypted PDFs are not supported")
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ResumeParseError:
        raise
    except Exception as exc:  # library exceptions vary across malformed PDFs
        raise ResumeParseError("invalid_pdf", f"Could not parse PDF: {exc}") from exc
    return "\n\n".join(pages), pages, "pypdf"


def _extract_docx(path: Path) -> tuple[str, list[str], str]:
    try:
        document = Document(path)
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    except Exception as exc:
        raise ResumeParseError("invalid_docx", f"Could not parse DOCX: {exc}") from exc
    return "\n".join(lines), [], "python-docx"


def _extract_txt(path: Path) -> tuple[str, list[str], str]:
    try:
        return path.read_text(encoding="utf-8", errors="replace"), [], "utf-8"
    except OSError as exc:
        raise ResumeParseError("unreadable_file", f"Could not read text file: {exc}") from exc


def parse_resume(path: Path, settings: Settings) -> ParsedResume:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise ResumeParseError("unreadable_file", f"Could not inspect file: {exc}") from exc
    if size > settings.max_resume_bytes:
        raise ResumeParseError("file_too_large", f"File exceeds {settings.max_resume_bytes} bytes")

    extension = path.suffix.lower()
    if extension == ".pdf":
        text, pages, parser = _extract_pdf(path)
    elif extension == ".docx":
        text, pages, parser = _extract_docx(path)
    elif extension == ".txt":
        text, pages, parser = _extract_txt(path)
    else:
        raise ResumeParseError("unsupported_type", f"Unsupported extension: {extension}")

    text, truncated = _normalise_text(text, settings.max_resume_characters)
    if not text:
        raise ResumeParseError("empty_resume", "No extractable text was found")
    digest = _sha256(path)
    warnings = ["Resume text was truncated to the configured limit."] if truncated else []
    return ParsedResume(
        file=ResumeFile(
            file_id=digest[:16],
            path=path,
            source_file=path.name,
            extension=extension,
            sha256=digest,
            size_bytes=size,
        ),
        raw_text=text,
        pages=pages,
        parser_name=parser,
        warnings=warnings,
    )
